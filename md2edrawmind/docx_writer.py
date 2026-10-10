# -*- coding: utf-8 -*-
"""生成 Word（.docx）：层级用 w:outlineLvl，公式是内嵌图片。

为什么是这种形态（都是实测结论）：
  - EdrawMind 的 Word 导入是**本地解析器**，不走云端 AI（不会把简体转繁体）
  - 它**不解析 OMML**（Word 原生公式会被拍平成文字），所以公式只能用图片承载
  - 层级必须写 w:outlineLvl，仅靠字体/缩进它不认
"""
from __future__ import annotations

import json
import os
import re
from dataclasses import dataclass
from typing import Dict, List, Tuple

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .parser import Node

FONT_PT = 10.5
PNG_DPI = 96.0          # 公式 PNG 以 16px 基准渲染
REF_PX = 16.0


@dataclass
class FormulaImage:
    path: str
    width: int
    height: int

    @property
    def width_pt(self) -> float:
        return self.width * 72.0 / PNG_DPI * (FONT_PT / REF_PX)

    @property
    def height_pt(self) -> float:
        return self.height * 72.0 / PNG_DPI * (FONT_PT / REF_PX)


def load_images(work_dir: str) -> Dict[str, FormulaImage]:
    """优先读 render.js 产出的 formula-pngs.json；退回 formulas.txt 序号约定。"""
    manifest = os.path.join(work_dir, "formula-pngs.json")
    out: Dict[str, FormulaImage] = {}
    if os.path.exists(manifest):
        for item in json.load(open(manifest, encoding="utf-8")):
            out[item["tex"]] = FormulaImage(item["file"], int(item["width"]), int(item["height"]))
        if out:
            return out
    formulas_path = os.path.join(work_dir, "formulas.txt")
    if os.path.exists(formulas_path):
        for i, line in enumerate(open(formulas_path, encoding="utf-8").read().splitlines(), 1):
            if "\t" not in line:
                continue
            tex = line.split("\t", 1)[1]
            p = os.path.join(work_dir, "formula-png", "f%03d.png" % i)
            if os.path.exists(p):
                from PIL import Image
                with Image.open(p) as im:
                    out[tex] = FormulaImage(p, im.width, im.height)
    return out


def set_outline_level(paragraph, level: int) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    for old in pPr.findall(qn("w:outlineLvl")):
        pPr.remove(old)
    el = OxmlElement("w:outlineLvl")
    el.set(qn("w:val"), str(max(0, min(8, level))))
    pPr.append(el)


def count_nodes(node: Node) -> int:
    return 1 + sum(count_nodes(c) for c in node.children)


def _style_run(run, node: Node) -> None:
    if node.kind == "heading":
        run.bold = True
        run.font.color.rgb = RGBColor(0x1A, 0x4F, 0x8A)
        run.font.size = Pt(FONT_PT + (1.5 if node.depth <= 2 else 0.5))


class DocxBuilder:
    def __init__(self, images: Dict[str, FormulaImage], title: str = ""):
        self.images = images
        self.title = title
        self.missing: List[str] = []
        self.embedded = 0

    def build(self, root: Node, out_path: str, part_title: str = "", nodes: List[Node] | None = None) -> dict:
        doc = Document()
        normal = doc.styles["Normal"]
        normal.font.name = "Microsoft YaHei"
        normal.font.size = Pt(FONT_PT)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

        head = part_title or self.title or "思维导图"
        p = doc.add_paragraph()
        run = p.add_run(head)
        run.bold = True
        run.font.size = Pt(FONT_PT + 6)

        for child in (nodes if nodes is not None else root.children):
            self._walk(doc, child)

        doc.save(out_path)
        return {"out": out_path, "nodes": sum(count_nodes(n) for n in (nodes or root.children)),
                "embedded": self.embedded, "missing": len(self.missing)}

    def _walk(self, doc, node: Node, base_depth: int = 0) -> None:
        para = doc.add_paragraph()
        pf = para.paragraph_format
        pf.space_before = Pt(2)
        pf.space_after = Pt(2)
        depth = max(0, node.depth - 1 + base_depth)
        pf.left_indent = Pt(depth * 14)
        set_outline_level(para, depth)

        for seg in node.segs:
            if seg.type == "text":
                text = re.sub(r"\s+", " ", seg.value)
                if not text.strip():
                    continue
                _style_run(para.add_run(text), node)
            else:
                img = self.images.get(seg.value)
                if img is None:
                    self.missing.append(seg.value)
                    para.add_run(seg.value)
                    continue
                para.add_run().add_picture(img.path, width=Pt(img.width_pt), height=Pt(img.height_pt))
                self.embedded += 1
        for child in node.children:
            self._walk(doc, child, base_depth)

    # ---------- 按主题数上限切分 ----------
    @staticmethod
    def split(root: Node, limit: int) -> List[Tuple[List[str], List[Node]]]:
        """把主题树切成若干份，每份不超过 limit 个节点。

        返回 [(祖先标题链, 节点列表), ...]。
        分支太大就往下钻（把子节点提到顶层），并用祖先标题链说明它原本属于谁。
        """
        units: List[Tuple[List[str], Node]] = []

        def walk(nodes: List[Node], prefix: List[str]) -> None:
            for n in nodes:
                size = count_nodes(n)
                if size <= limit:
                    units.append((prefix, n))          # 整支作为一份
                else:
                    # 太大：自己（标题）算一个主题，子节点提升，用标题链说明归属
                    units.append((prefix, Node(n.kind, n.depth, list(n.segs), [])))
                    if n.children:
                        walk(n.children, prefix + [n.text or n.kind])

        walk(root.children, [])

        parts: List[Tuple[List[str], List[Node]]] = []
        cur: List[Node] = []
        cur_prefix: List[str] = []
        used = 0
        for prefix, node in units:
            size = count_nodes(node)
            if cur and (used + size > limit or prefix != cur_prefix):
                parts.append((cur_prefix, cur))
                cur, used, cur_prefix = [], 0, []
            if not cur:
                cur_prefix = prefix
            cur.append(node)
            used += size
        if cur:
            parts.append((cur_prefix, cur))
        return parts
